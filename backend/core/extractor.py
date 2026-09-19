import os
from PyPDF2 import PdfReader
import docx
from pptx import Presentation

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    text = ""

    if ext == ".pdf":
        reader = PdfReader(filepath)
        for page in reader.pages:
            text += page.extract_text() or ""

    elif ext == ".docx":
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            text += para.text + "\n"

    elif ext == ".pptx":
        prs = Presentation(filepath)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            text += run.text + " "

    elif ext == ".txt":
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return text